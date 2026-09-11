from __future__ import annotations

import re
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from typing import Any
from uuid import uuid4
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ddbox_api.domain.errors import ValidationError
from ddbox_api.domain.models import AttributionTouch, StoredLead, utc_now
from ddbox_api.repositories.base import ContentRepository

EXPORT_SCHEMA_VERSION = 1
EXPORT_PAGE_SIZE = 500
_BANGKOK = ZoneInfo("Asia/Bangkok")
_ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA_PREFIX = re.compile(r"^\s*[=+\-@]")

LEAD_EXPORT_COLUMNS = (
    "lead_id",
    "lead_reference",
    "created_at_utc",
    "received_at_th",
    "lead_origin",
    "contact_name",
    "company",
    "preferred_contact",
    "phone",
    "line_id",
    "email",
    "customer_path",
    "product_type",
    "quantity",
    "delivery_province",
    "required_date",
    "project_details",
    "submission_path",
    "attribution_status",
    "attribution_drop_reason",
    "measurement_consent_mode",
    "measurement_consent_version",
    "measurement_consent_updated_at_utc",
    "first_utm_source",
    "first_utm_medium",
    "first_utm_campaign",
    "first_utm_id",
    "first_utm_content",
    "first_utm_term",
    "first_adgroup_id",
    "first_landing_path",
    "first_captured_at_utc",
    "last_utm_source",
    "last_utm_medium",
    "last_utm_campaign",
    "last_utm_id",
    "last_utm_content",
    "last_utm_term",
    "last_adgroup_id",
    "last_landing_path",
    "last_captured_at_utc",
    "gclid_present",
    "source_environment",
    "exported_at_utc",
    "export_id",
    "export_schema_version",
)


@dataclass(frozen=True)
class LeadExport:
    content: bytes
    export_id: str
    filename: str
    record_count: int
    created_at: datetime


def _excel_datetime(value: datetime, timezone: ZoneInfo | None = None) -> datetime:
    target = value.astimezone(timezone or UTC)
    return target.replace(tzinfo=None)


def _touch_value(touch: AttributionTouch | None, field: str) -> Any:
    return getattr(touch, field) if touch is not None else None


def _safe_excel_text(value: str) -> str:
    sanitized = _ILLEGAL_EXCEL_CHARACTERS.sub(" ", value)
    return f"'{sanitized}" if _FORMULA_PREFIX.match(sanitized) else sanitized


def _cell(sheet: Any, value: Any, *, header: bool = False) -> Any:
    if isinstance(value, str):
        value = _safe_excel_text(value)
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
        cell.number_format = "@"
    elif isinstance(value, datetime):
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    if header:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="111827")
    return cell


class LeadExportService:
    def __init__(
        self,
        repository: ContentRepository,
        environment: str,
        *,
        page_size: int = EXPORT_PAGE_SIZE,
    ) -> None:
        self._repository = repository
        self._environment = environment
        self._page_size = page_size

    def _read_leads(self, created_from: datetime, created_to: datetime) -> Iterator[StoredLead]:
        after_id: str | None = None
        while True:
            page = self._repository.list_leads(
                created_from,
                created_to,
                after_id=after_id,
                limit=self._page_size,
            )
            if not page:
                return
            yield from page
            next_after_id = page[-1].id
            if next_after_id == after_id:
                raise RuntimeError("lead export pagination cursor did not advance")
            after_id = next_after_id
            if len(page) < self._page_size:
                return

    def _lead_row(
        self,
        lead: StoredLead,
        *,
        export_id: str,
        exported_at: datetime,
    ) -> list[Any]:
        payload = lead.payload
        attribution = payload.attribution
        first = attribution.first_touch if attribution else None
        last = attribution.last_touch if attribution else None
        consent = payload.measurement_consent
        values: dict[str, Any] = {
            "lead_id": lead.id,
            "lead_reference": lead.reference,
            "created_at_utc": _excel_datetime(lead.created_at),
            "received_at_th": _excel_datetime(lead.created_at, _BANGKOK),
            "lead_origin": lead.lead_origin,
            "contact_name": payload.contact_name,
            "company": payload.company,
            "preferred_contact": payload.preferred_contact.value,
            "phone": payload.phone,
            "line_id": payload.line_id,
            "email": str(payload.email) if payload.email is not None else None,
            "customer_path": payload.customer_path.value,
            "product_type": payload.product_type,
            "quantity": payload.quantity,
            "delivery_province": payload.delivery_province,
            "required_date": payload.required_date,
            "project_details": payload.project_details,
            "submission_path": payload.submission_path,
            "attribution_status": lead.attribution_status.value,
            "attribution_drop_reason": (
                lead.attribution_drop_reason.value if lead.attribution_drop_reason else None
            ),
            "measurement_consent_mode": consent.mode.value if consent else None,
            "measurement_consent_version": consent.version if consent else None,
            "measurement_consent_updated_at_utc": (
                _excel_datetime(consent.updated_at) if consent and consent.updated_at else None
            ),
            "first_utm_source": _touch_value(first, "utm_source"),
            "first_utm_medium": _touch_value(first, "utm_medium"),
            "first_utm_campaign": _touch_value(first, "utm_campaign"),
            "first_utm_id": _touch_value(first, "utm_id"),
            "first_utm_content": _touch_value(first, "utm_content"),
            "first_utm_term": _touch_value(first, "utm_term"),
            "first_adgroup_id": _touch_value(first, "adgroup_id"),
            "first_landing_path": _touch_value(first, "landing_path"),
            "first_captured_at_utc": (_excel_datetime(first.captured_at) if first else None),
            "last_utm_source": _touch_value(last, "utm_source"),
            "last_utm_medium": _touch_value(last, "utm_medium"),
            "last_utm_campaign": _touch_value(last, "utm_campaign"),
            "last_utm_id": _touch_value(last, "utm_id"),
            "last_utm_content": _touch_value(last, "utm_content"),
            "last_utm_term": _touch_value(last, "utm_term"),
            "last_adgroup_id": _touch_value(last, "adgroup_id"),
            "last_landing_path": _touch_value(last, "landing_path"),
            "last_captured_at_utc": _excel_datetime(last.captured_at) if last else None,
            "gclid_present": "Yes" if (first and first.gclid) or (last and last.gclid) else "No",
            "source_environment": self._environment,
            "exported_at_utc": _excel_datetime(exported_at),
            "export_id": export_id,
            "export_schema_version": EXPORT_SCHEMA_VERSION,
        }
        return [values[column] for column in LEAD_EXPORT_COLUMNS]

    def create(self, created_from: datetime, created_to: datetime) -> LeadExport:
        if (
            created_from.tzinfo is None
            or created_from.utcoffset() is None
            or created_to.tzinfo is None
            or created_to.utcoffset() is None
        ):
            raise ValidationError("export timestamps must include a timezone")
        if created_to <= created_from:
            raise ValidationError("created_to must be after created_from")
        exported_at = utc_now()
        snapshot_cutoff = min(created_to, exported_at)
        if snapshot_cutoff <= created_from:
            raise ValidationError("created_from must be before the export cutoff")
        export_id = uuid4().hex
        workbook = Workbook(write_only=True)
        workbook.iso_dates = True

        lead_sheet = workbook.create_sheet("Leads")
        lead_sheet.freeze_panes = "A2"
        lead_sheet.append(
            [_cell(lead_sheet, column, header=True) for column in LEAD_EXPORT_COLUMNS]
        )
        record_count = 0
        for lead in self._read_leads(created_from, snapshot_cutoff):
            lead_sheet.append(
                [
                    _cell(lead_sheet, value)
                    for value in self._lead_row(
                        lead,
                        export_id=export_id,
                        exported_at=exported_at,
                    )
                ]
            )
            record_count += 1
        lead_sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(LEAD_EXPORT_COLUMNS))}{record_count + 1}"
        )

        metadata_sheet = workbook.create_sheet("Export_Metadata")
        metadata_sheet.append(
            [
                _cell(metadata_sheet, "field", header=True),
                _cell(metadata_sheet, "value", header=True),
            ]
        )
        metadata = (
            ("export_id", export_id),
            ("export_schema_version", EXPORT_SCHEMA_VERSION),
            ("created_at_utc", _excel_datetime(exported_at)),
            ("source_environment", self._environment),
            ("range_start_utc", _excel_datetime(created_from)),
            ("range_end_utc_exclusive", _excel_datetime(snapshot_cutoff)),
            ("record_count", record_count),
            ("attribution_model", "first_last_tagged"),
            ("timezone_note", "UTC fields are UTC; received_at_th is Asia/Bangkok"),
        )
        for key, value in metadata:
            metadata_sheet.append([_cell(metadata_sheet, key), _cell(metadata_sheet, value)])

        output = BytesIO()
        workbook.save(output)
        filename_timestamp = exported_at.strftime("%Y%m%dT%H%M%SZ")
        return LeadExport(
            content=output.getvalue(),
            export_id=export_id,
            filename=f"DD_BOX_Leads_{filename_timestamp}.xlsx",
            record_count=record_count,
            created_at=exported_at,
        )
