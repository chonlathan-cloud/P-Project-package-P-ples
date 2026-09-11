from __future__ import annotations

from datetime import UTC, datetime, timedelta
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from ddbox_api.services.lead_exports import LeadExportService


def _lead_payload() -> dict[str, object]:
    captured_at = datetime.now(UTC) - timedelta(hours=1)
    touch = {
        "utm_source": "google",
        "utm_medium": "cpc",
        "utm_campaign": "dd45_core_custom_box",
        "utm_id": "1111111111",
        "utm_content": "3333333333",
        "utm_term": "รับทำกล่องออฟเซ็ท",
        "adgroup_id": "2222222222",
        "gclid": "SECRET-GCLID-NOT-FOR-SALES",
        "landing_path": "/products/folding-carton",
        "captured_at": captured_at.isoformat(),
        "expires_at": (captured_at + timedelta(days=90)).isoformat(),
    }
    return {
        "customer_path": "has_specifications",
        "product_type": "=2+2",
        "quantity": 1000,
        "project_details": "ต้องการประเมินรูปแบบกล่องและวัสดุสำหรับสินค้าใหม่",
        "contact_name": "ผู้ติดต่อทดสอบ",
        "phone": "+66812345678",
        "preferred_contact": "phone",
        "consent": True,
        "submission_path": "/quote",
        "measurement_consent": {
            "mode": "all",
            "version": 1,
            "updated_at": (captured_at - timedelta(minutes=1)).isoformat(),
        },
        "attribution": {
            "schema_version": 1,
            "model": "first_last_tagged",
            "first_touch": touch,
            "last_touch": touch,
        },
    }


def _range_parameters() -> dict[str, str]:
    now = datetime.now(UTC)
    return {
        "created_from": (now - timedelta(days=1)).isoformat(),
        "created_to": (now + timedelta(days=1)).isoformat(),
    }


def test_lead_export_requires_admin(client: TestClient) -> None:
    response = client.get("/v1/admin/leads/export", params=_range_parameters())

    assert response.status_code == 403


def test_lead_export_is_a_safe_read_only_snapshot(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    created = client.post(
        "/v1/leads",
        headers={"Idempotency-Key": "stable-client-key-export"},
        json=_lead_payload(),
    )
    assert created.status_code == 201

    response = client.get(
        "/v1/admin/leads/export",
        headers=admin_headers,
        params=_range_parameters(),
    )

    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["x-ddbox-record-count"] == "1"
    assert response.headers["x-ddbox-export-id"]
    assert "attachment;" in response.headers["content-disposition"]
    assert b"SECRET-GCLID-NOT-FOR-SALES" not in response.content

    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
    assert workbook.sheetnames == ["Leads", "Export_Metadata"]
    sheet = workbook["Leads"]
    headers = [cell.value for cell in sheet[1]]
    values = [cell.value for cell in sheet[2]]
    row = dict(zip(headers, values, strict=True))

    assert "gclid" not in headers
    assert "SECRET-GCLID-NOT-FOR-SALES" not in values
    assert row["gclid_present"] == "Yes"
    assert row["first_utm_id"] == "1111111111"
    assert row["last_adgroup_id"] == "2222222222"
    assert row["product_type"] == "'=2+2"
    assert sheet.cell(row=2, column=headers.index("product_type") + 1).data_type == "s"
    assert row["phone"] == "'+66812345678"
    assert isinstance(row["created_at_utc"], datetime)
    assert isinstance(row["received_at_th"], datetime)

    metadata = {
        row[0].value: row[1].value for row in workbook["Export_Metadata"].iter_rows(min_row=2)
    }
    assert metadata["record_count"] == 1
    assert metadata["source_environment"] == "test"
    assert metadata["export_id"] == response.headers["x-ddbox-export-id"]
    assert metadata["range_end_utc_exclusive"] <= datetime.now()


def test_lead_export_paginates_without_dropping_records(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    client.app.state.lead_export_service = LeadExportService(
        client.app.state.repository,
        "test",
        page_size=1,
    )
    for suffix in ("one", "two"):
        created = client.post(
            "/v1/leads",
            headers={"Idempotency-Key": f"stable-client-key-export-{suffix}"},
            json=_lead_payload(),
        )
        assert created.status_code == 201

    response = client.get(
        "/v1/admin/leads/export",
        headers=admin_headers,
        params=_range_parameters(),
    )

    assert response.status_code == 200
    assert response.headers["x-ddbox-record-count"] == "2"
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert sum(1 for _row in workbook["Leads"].iter_rows()) == 3


def test_lead_export_rejects_an_invalid_range(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    now = datetime.now(UTC)
    reversed_range = client.get(
        "/v1/admin/leads/export",
        headers=admin_headers,
        params={
            "created_from": now.isoformat(),
            "created_to": (now - timedelta(days=1)).isoformat(),
        },
    )
    assert reversed_range.status_code == 422
